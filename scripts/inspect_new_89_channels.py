import openpyxl
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('available_channels_for_review.xlsx')
ws = wb.active

new_channels = []
for r in range(2, ws.max_row + 1):
    genre = ws.cell(r, 2).value
    name = ws.cell(r, 3).value
    quality = ws.cell(r, 4).value
    latency = ws.cell(r, 5).value
    bdix = ws.cell(r, 6).value
    feeds = ws.cell(r, 7).value
    note = ws.cell(r, 8).value
    url = ws.cell(r, 9).value
    logo = ws.cell(r, 10).value
    if name and str(name).strip():
        new_channels.append({
            'row': r,
            'genre': str(genre).strip() if genre else '',
            'name': str(name).strip(),
            'quality': quality,
            'latency': latency,
            'bdix': bdix,
            'feeds': feeds,
            'url': str(url).strip() if url else '',
            'logo': str(logo).strip() if logo else ''
        })

print(f"Total new channels to add: {len(new_channels)}")
for ch in new_channels:
    print(f"{ch['row']:2}. {ch['name']:<30} | {ch['genre']:<30} | Logo: {ch['logo'][:40]} | URL: {ch['url'][:50]}")
