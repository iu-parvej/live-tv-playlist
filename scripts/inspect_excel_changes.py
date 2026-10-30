import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

# 1. Read kept curated channels
wb_curated = openpyxl.load_workbook('curated_channels.xlsx')
ws_curated = wb_curated.active

kept_curated = []
for r in range(2, ws_curated.max_row + 1):
    sl = ws_curated.cell(r, 1).value
    cat = ws_curated.cell(r, 2).value
    name = ws_curated.cell(r, 3).value
    streams = ws_curated.cell(r, 4).value
    logo = ws_curated.cell(r, 5).value
    if name and str(name).strip():
        kept_curated.append({
            'orig_sl': sl,
            'category': str(cat).strip() if cat else '',
            'name': str(name).strip(),
            'streams': streams,
            'logo': logo
        })

print(f"Total kept channels in curated_channels.xlsx: {len(kept_curated)}")

# 2. Read kept review channels
wb_review = openpyxl.load_workbook('available_channels_for_review.xlsx')
ws_review = wb_review.active

kept_review = []
for r in range(2, ws_review.max_row + 1):
    sl = ws_review.cell(r, 1).value
    genre = ws_review.cell(r, 2).value
    name = ws_review.cell(r, 3).value
    quality = ws_review.cell(r, 4).value
    latency = ws_review.cell(r, 5).value
    bdix = ws_review.cell(r, 6).value
    feeds = ws_review.cell(r, 7).value
    note = ws_review.cell(r, 8).value
    url = ws_review.cell(r, 9).value
    logo = ws_review.cell(r, 10).value
    if name and str(name).strip():
        kept_review.append({
            'orig_sl': sl,
            'genre': str(genre).strip() if genre else '',
            'name': str(name).strip(),
            'quality': quality,
            'latency': latency,
            'bdix': bdix,
            'feeds': feeds,
            'note': note,
            'url': url,
            'logo': logo
        })

print(f"Total channels to add from available_channels_for_review.xlsx: {len(kept_review)}")
print("\n--- SAMPLE CHANNELS TO ADD (First 20) ---")
for i, ch in enumerate(kept_review[:20], 1):
    print(f"{i:2}. [{ch['genre']}] {ch['name']} (Latency: {ch['latency']}, BDIX: {ch['bdix']})")

if len(kept_review) > 20:
    print(f"\n--- SAMPLE CHANNELS TO ADD (21 to {min(40, len(kept_review))}) ---")
    for i, ch in enumerate(kept_review[20:40], 21):
        print(f"{i:2}. [{ch['genre']}] {ch['name']} (Latency: {ch['latency']}, BDIX: {ch['bdix']})")
