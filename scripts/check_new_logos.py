import openpyxl
import os
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('available_channels_for_review.xlsx')
ws = wb.active

local_logos = set(os.listdir('logos'))

channels = []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 3).value
    logo = ws.cell(r, 10).value
    if name and str(name).strip():
        channels.append({'name': str(name).strip(), 'logo': str(logo).strip() if logo else ''})

print(f"Total review channels: {len(channels)}")

missing_local = []
for ch in channels:
    slug = re.sub(r'[^a-z0-9]+', '_', ch['name'].lower()).strip('_')
    slug_no_hd = re.sub(r'_(?:hd|fhd|sd|4k|tv)$', '', slug)
    
    found = False
    for ext in ['png', 'jpg']:
        if f"{slug}.{ext}" in local_logos or f"{slug_no_hd}.{ext}" in local_logos:
            found = True
            break
    if not found:
        missing_local.append((ch['name'], slug, ch['logo']))

print(f"Channels with existing local logo: {len(channels) - len(missing_local)}")
print(f"Channels needing logo download: {len(missing_local)}")
for name, slug, ext_logo in missing_local:
    print(f"  Need logo: {name:<28} | slug: {slug:<25} | External: {ext_logo[:50]}")
