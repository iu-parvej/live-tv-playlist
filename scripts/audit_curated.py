import os
import sys
import re
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from core.normalizer import ChannelNormalizer

# 1. Parse playlist_curated.m3u
curated_file = "playlist_curated.m3u"
entries = []
categories = {}
logos = {}

with open(curated_file, "r", encoding="utf-8") as f:
    curr = {}
    for line in f:
        line = line.strip()
        if line.startswith("#EXTINF:"):
            curr = {}
            m_logo = re.search(r'tvg-logo="([^"]+)"', line)
            m_grp = re.search(r'group-title="([^"]+)"', line)
            m_id = re.search(r'tvg-id="([^"]+)"', line)
            title = line.split(",")[-1].strip()
            curr["name"] = title
            curr["logo"] = m_logo.group(1) if m_logo else ""
            curr["group"] = m_grp.group(1) if m_grp else ""
            curr["tvg_id"] = m_id.group(1) if m_id else ""
        elif line.startswith("http://") or line.startswith("https://"):
            if curr:
                curr["url"] = line
                entries.append(curr)
                cat = curr["group"]
                categories[cat] = categories.get(cat, 0) + 1
                curr = {}

print("=" * 65)
print(f" AUDIT REPORT: {curated_file}")
print("=" * 65)
print(f"Total streams in curated playlist: {len(entries)}")

# Distinct channel names (strip numbering like ' 1', ' 2', ' 3')
distinct_channels = set()
for e in entries:
    base_name = re.sub(r'\s+\d+$', '', e['name']).strip()
    distinct_channels.add((base_name, e['group']))

print(f"Total distinct channels: {len(distinct_channels)}")

print("\n--- CATEGORY BREAKDOWN ---")
for cat, count in sorted(categories.items()):
    # Count distinct channels in this category
    chans_in_cat = set(re.sub(r'\s+\d+$', '', e['name']).strip() for e in entries if e['group'] == cat)
    print(f"  {cat:<45} : {len(chans_in_cat):2} channels ({count:3} feeds)")

# 2. Check Blacklist Leaks
with open("config/curated_blacklist.json", "r", encoding="utf-8") as f:
    bl_data = json.load(f)

leaked = []
for e in entries:
    n = e['name'].lower()
    for b in bl_data.get("blacklist", []):
        b_slug = re.sub(r'[^a-z0-9]+', '', b.lower())
        n_slug = re.sub(r'[^a-z0-9]+', '', n)
        if b_slug in n_slug or b.lower() in n:
            leaked.append((e['name'], b, e['url']))

print("\n--- BLACKLIST CHECK (33 Unwanted Channels) ---")
if leaked:
    print(f"  WARNING: {len(leaked)} blacklisted channel streams found!")
    for name, bl_name, url in leaked:
        print(f"    LEAKED: '{name}' matched blacklist '{bl_name}'")
else:
    print("  [PASSED] 0 Blacklisted channels found in curated playlist (100% CLEAN)")

# 3. Check Logo Coverage
missing_logo = []
non_github_logo = []
for e in entries:
    logo = e['logo']
    if not logo:
        missing_logo.append(e['name'])
    elif "raw.githubusercontent.com/iu-parvej/live-tv-playlist" not in logo:
        non_github_logo.append((e['name'], logo))

print("\n--- LOGO INTEGRITY CHECK ---")
print(f"  Missing logos: {len(missing_logo)}")
print(f"  External non-GitHub logos: {len(set(x[0] for x in non_github_logo))}")
if non_github_logo:
    for name, logo in list(set(non_github_logo))[:10]:
        print(f"    External: {name:<25} -> {logo[:60]}")

# 4. Check 89 Review Channels
norm = ChannelNormalizer()

wb_review = openpyxl.load_workbook("available_channels_for_review.xlsx")
ws_review = wb_review.active
review_channels = []
for r in range(2, ws_review.max_row + 1):
    n = ws_review.cell(r, 3).value
    if n and str(n).strip():
        review_channels.append(str(n).strip())

curated_base_names = set(re.sub(r'\s+\d+$', '', e['name']).strip().lower() for e in entries)

missing_review = []
for rc in review_channels:
    clean_rc = norm.clean_name(rc).lower()
    rc_norm = re.sub(r'[^a-z0-9]+', '', rc.lower())
    clean_norm = re.sub(r'[^a-z0-9]+', '', clean_rc)
    
    found = False
    for cb in curated_base_names:
        cb_norm = re.sub(r'[^a-z0-9]+', '', cb.lower())
        if clean_rc == cb or clean_norm == cb_norm or rc_norm == cb_norm or rc.lower() in cb or cb in rc.lower() or clean_rc in cb or cb in clean_rc:
            found = True
            break
    if not found:
        missing_review.append(rc)

print(f"\n--- 89 REVIEW CHANNELS INCLUSION ---")
print(f"  Total review channels to include: {len(review_channels)}")
print(f"  Successfully found in curated playlist: {len(review_channels) - len(missing_review)}")
print(f"  Missing: {len(missing_review)}")
for m in missing_review:
    print(f"    Missing: {m}")

print("\n" + "=" * 65)
