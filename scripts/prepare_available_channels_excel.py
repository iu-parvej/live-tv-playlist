import re
import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

# 1. Load curated channel names to exclude them
curated_names = set()
with open('playlist_curated.m3u', 'r', encoding='utf-8') as f:
    for line in f:
        if line.startswith('#EXTINF:'):
            full_name = line.split(',')[-1].strip()
            base_name = re.sub(r'\s+\d+$', '', full_name).strip()
            slug = re.sub(r'[^a-z0-9]+', '', base_name.lower())
            curated_names.add(slug)

print(f"Curated channels already present: {len(curated_names)}")

# 2. Load all working channels from output/playlist_multiserver.m3u
# This gives us all streams, latency, bdix status, and logo per channel
streams_by_channel = {}
with open('output/playlist_multiserver.m3u', 'r', encoding='utf-8') as f:
    curr = {}
    for line in f:
        line = line.strip()
        if line.startswith('#EXTINF:'):
            curr = {}
            m_logo = re.search(r'tvg-logo="([^"]+)"', line)
            m_grp = re.search(r'group-title="([^"]+)"', line)
            title = line.split(',')[-1].strip()
            curr['raw_name'] = title
            curr['logo'] = m_logo.group(1) if m_logo else ''
            curr['group'] = m_grp.group(1) if m_grp else 'General'
        elif line.startswith('http://') or line.startswith('https://'):
            if curr:
                curr['url'] = line
                # Extract latency & BDIX from title if present
                m_lat = re.search(r'\(.*?(?:Server \d+ - )?(\d+)ms.*?\)', curr['raw_name'])
                curr['latency_ms'] = int(m_lat.group(1)) if m_lat else 999
                curr['is_bdix'] = '[BDIX]' in curr['raw_name']
                # Clean base channel name
                c_name = re.sub(r'\s*\[BDIX\]\s*', '', curr['raw_name'], flags=re.I)
                c_name = re.sub(r'\s*\((?:Server \d+ - )?\d+ms.*?\)', '', c_name)
                c_name = re.sub(r'\s+\d+$', '', c_name).strip()
                
                slug = re.sub(r'[^a-z0-9]+', '', c_name.lower())
                if slug not in streams_by_channel:
                    streams_by_channel[slug] = {
                        'name': c_name,
                        'group': curr['group'],
                        'logo': curr['logo'],
                        'streams': []
                    }
                streams_by_channel[slug]['streams'].append(curr)
                curr = {}

print(f"Total distinct active channels across all playlists: {len(streams_by_channel)}")

# 3. Filter out channels that are ALREADY in Curated VIP
candidates = []
for slug, ch in streams_by_channel.items():
    if slug in curated_names:
        continue
    # If partial match or alias matches curated
    if any(c in slug for c in curated_names if len(c) > 6):
        continue
    
    # Sort streams by latency
    ch['streams'].sort(key=lambda x: x.get('latency_ms', 9999))
    best_stream = ch['streams'][0]
    best_latency = best_stream.get('latency_ms', 999)
    is_bdix = any(s.get('is_bdix', False) for s in ch['streams'])
    
    # Extract resolution if mentioned
    m_res = re.search(r'(1080p|720p|576p|480p|4k|hd|fhd)', ch['name'].lower())
    quality = m_res.group(1).upper() if m_res else ("1080p" if "1080" in best_stream['url'] else "HD/SD")
    
    # Categorize into user-friendly review genres
    nl = ch['name'].lower()
    gl = ch['group'].lower()
    
    genre = "General Entertainment"
    note = ""
    
    if any(w in nl or w in gl for w in ['sport', 'cricket', 'football', 'fifa', 'kabaddi', 'wwe']):
        genre = "⚽ Sports & Live Events"
        note = "Live Sports Stream"
    elif any(w in nl or w in gl for w in ['cartoon', 'kids', 'pogo', 'disney', 'nick', 'chithiram', 'yay', 'jungle']):
        genre = "👶 Kids & Cartoons"
        note = "Animation & Children"
    elif any(w in nl or w in gl for w in ['movie', 'cinema', 'film', 'max', 'goldmines', 'box office', 'bhojpuri cinema']):
        genre = "🎬 Movies & Cinema"
        note = "Hindi / South Dubbed / Movies"
    elif any(w in nl or w in gl for w in ['discovery', 'nat geo', 'history', 'wild', 'travel', 'science', 'animal', 'docu']):
        genre = "🌍 Documentary & Infotainment"
        note = "Wildlife & Documentaries"
    elif any(w in nl or w in gl for w in ['music', 'beats', '9x', 'mtv', 'sangeet', 'gaan', 'mastiii', 'hits', 'b4u music']):
        genre = "🎵 Music & Songs"
        note = "24x7 Music Channels"
    elif any(w in nl or w in gl for w in ['bangla', 'kolkata', 'jalsha', 'calcutta', 'rupashi', 'aakash']):
        genre = "🇧🇩 Kolkata & Bengali Regional"
        note = "Bengali Entertainment / News"
    elif any(w in nl or w in gl for w in ['bbc', 'cnn', 'al jazeera', 'dw', 'france 24', 'bloomberg', 'cnbc', 'wion', 'sky news', 'euronews']):
        genre = "🌐 International News (English)"
        note = "Top Global News Network"
    elif any(w in nl or w in gl for w in ['aaj tak', 'abp', 'ndtv', 'zee news', 'republic', 'india today', 'news18', 'tv9', 'sansad']):
        genre = "🇮🇳 Indian News (Hindi / National)"
        note = "Leading Indian News"
    elif any(w in nl or w in gl for w in ['quran', 'islam', 'makkah', 'madinah', 'sunnah', 'peace tv', 'madani', 'religious', 'bhakti', 'darshan']):
        genre = "🕌 Religious & Devotional"
        note = "Religious Broadcast"
    elif any(w in nl or w in gl for w in ['tamil', 'telugu', 'kannada', 'malayalam', 'marathi', 'punjabi', 'bhojpuri']):
        genre = "🇮🇳 Indian Regional Channels"
        note = "Regional Indian Language"
    
    candidates.append({
        'name': ch['name'],
        'genre': genre,
        'quality': quality,
        'latency_ms': best_latency,
        'is_bdix': is_bdix,
        'stream_count': len(ch['streams']),
        'best_url': best_stream['url'],
        'logo': ch['logo'],
        'note': note
    })

print(f"Total available candidate channels for review: {len(candidates)}")

# Order candidates by Genre priority, then latency
genre_order = [
    "⚽ Sports & Live Events",
    "👶 Kids & Cartoons",
    "🎬 Movies & Cinema",
    "🌍 Documentary & Infotainment",
    "🇧🇩 Kolkata & Bengali Regional",
    "🎵 Music & Songs",
    "🌐 International News (English)",
    "🇮🇳 Indian News (Hindi / National)",
    "🕌 Religious & Devotional",
    "General Entertainment",
    "🇮🇳 Indian Regional Channels"
]

def sort_key(c):
    g_idx = genre_order.index(c['genre']) if c['genre'] in genre_order else 99
    return (g_idx, c['latency_ms'])

candidates.sort(key=sort_key)

# Build Excel workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Available Channels"

header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
bdix_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Light Green

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

headers = [
    "SL",
    "Genre / Category",
    "Channel Name",
    "Resolution",
    "Best Latency",
    "BDIX Fast",
    "Total Feeds",
    "Recommendation Note",
    "Stream Sample URL",
    "Logo URL"
]
ws.append(headers)

for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border
ws.row_dimensions[1].height = 28

for sl, ch in enumerate(candidates, 1):
    row_idx = sl + 1
    row_data = [
        sl,
        ch['genre'],
        ch['name'],
        ch['quality'],
        f"{ch['latency_ms']} ms",
        "YES [BDIX]" if ch['is_bdix'] else "Standard",
        ch['stream_count'],
        ch['note'],
        ch['best_url'],
        ch['logo']
    ]
    ws.append(row_data)

    fill = bdix_fill if ch['is_bdix'] else (zebra_fill if sl % 2 == 0 else white_fill)

    # Serial
    c_sl = ws.cell(row=row_idx, column=1)
    c_sl.alignment = Alignment(horizontal="center", vertical="center")
    c_sl.font = Font(name="Segoe UI", size=10, bold=True)
    c_sl.fill = fill
    c_sl.border = thin_border

    # Genre
    c_g = ws.cell(row=row_idx, column=2)
    c_g.alignment = Alignment(horizontal="left", vertical="center")
    c_g.font = Font(name="Segoe UI", size=10, bold=True)
    c_g.fill = fill
    c_g.border = thin_border

    # Channel Name
    c_n = ws.cell(row=row_idx, column=3)
    c_n.alignment = Alignment(horizontal="left", vertical="center")
    c_n.font = Font(name="Segoe UI", size=10, bold=True)
    c_n.fill = fill
    c_n.border = thin_border

    # Quality
    c_q = ws.cell(row=row_idx, column=4)
    c_q.alignment = Alignment(horizontal="center", vertical="center")
    c_q.font = Font(name="Segoe UI", size=10)
    c_q.fill = fill
    c_q.border = thin_border

    # Latency
    c_lat = ws.cell(row=row_idx, column=5)
    c_lat.alignment = Alignment(horizontal="center", vertical="center")
    c_lat.font = Font(name="Segoe UI", size=10)
    c_lat.fill = fill
    c_lat.border = thin_border

    # BDIX
    c_bdix = ws.cell(row=row_idx, column=6)
    c_bdix.alignment = Alignment(horizontal="center", vertical="center")
    c_bdix.font = Font(name="Segoe UI", size=10, bold=ch['is_bdix'])
    c_bdix.fill = fill
    c_bdix.border = thin_border

    # Feeds count
    c_fc = ws.cell(row=row_idx, column=7)
    c_fc.alignment = Alignment(horizontal="center", vertical="center")
    c_fc.font = Font(name="Segoe UI", size=10)
    c_fc.fill = fill
    c_fc.border = thin_border

    # Note
    c_nt = ws.cell(row=row_idx, column=8)
    c_nt.alignment = Alignment(horizontal="left", vertical="center")
    c_nt.font = Font(name="Segoe UI", size=9, italic=True)
    c_nt.fill = fill
    c_nt.border = thin_border

    # URL
    c_u = ws.cell(row=row_idx, column=9)
    c_u.alignment = Alignment(horizontal="left", vertical="center")
    c_u.font = Font(name="Segoe UI", size=9, color="595959")
    c_u.fill = fill
    c_u.border = thin_border

    # Logo
    c_l = ws.cell(row=row_idx, column=10)
    c_l.alignment = Alignment(horizontal="left", vertical="center")
    c_l.font = Font(name="Segoe UI", size=9, color="595959")
    c_l.fill = fill
    c_l.border = thin_border

    ws.row_dimensions[row_idx].height = 22

# Column dimensions
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 34
ws.column_dimensions['D'].width = 14
ws.column_dimensions['E'].width = 14
ws.column_dimensions['F'].width = 14
ws.column_dimensions['G'].width = 12
ws.column_dimensions['H'].width = 28
ws.column_dimensions['I'].width = 45
ws.column_dimensions['J'].width = 40

ws.freeze_panes = "A2"

excel_path = "available_channels_for_review.xlsx"
wb.save(excel_path)
print(f"Saved review list to: {excel_path}")

# Also CSV
import csv
with open('available_channels_for_review.csv', 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    for row in ws.iter_rows(values_only=True):
        writer.writerow(row)
print("Saved review list to: available_channels_for_review.csv")
