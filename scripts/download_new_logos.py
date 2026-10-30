import openpyxl
import os
import sys
import re
import asyncio
import aiohttp

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('available_channels_for_review.xlsx')
ws = wb.active

local_logos = set(os.listdir('logos'))

download_tasks = []
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 3).value
    logo = ws.cell(r, 10).value
    if name and str(name).strip():
        ch_name = str(name).strip()
        slug = re.sub(r'[^a-z0-9]+', '_', ch_name.lower()).strip('_')
        slug_no_hd = re.sub(r'_(?:hd|fhd|sd|4k|tv)$', '', slug)
        
        target_file = f"{slug}.png"
        if target_file not in local_logos and f"{slug_no_hd}.png" not in local_logos:
            logo_url = str(logo).strip() if logo else ''
            if logo_url and logo_url.startswith('http'):
                download_tasks.append((ch_name, target_file, logo_url))

print(f"Total logos to download: {len(download_tasks)}")

async def download_logo(session, sem, name, filename, url):
    target_path = os.path.join('logos', filename)
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
    }
    async with sem:
        try:
            async with session.get(url, headers=headers, timeout=aiohttp.ClientTimeout(total=8), ssl=False) as resp:
                if resp.status == 200:
                    content = await resp.read()
                    if len(content) > 300:
                        with open(target_path, 'wb') as f:
                            f.write(content)
                        return True, name, len(content)
        except Exception as e:
            pass
    return False, name, 0

async def main():
    sem = asyncio.Semaphore(15)
    async with aiohttp.ClientSession() as session:
        coros = [download_logo(session, sem, name, fn, url) for name, fn, url in download_tasks]
        results = await asyncio.gather(*coros)
    
    success = [r for r in results if r[0]]
    failed = [r for r in results if not r[0]]
    print(f"Successfully downloaded: {len(success)} logos")
    if failed:
        print(f"Failed to download: {len(failed)} logos")
        for f in failed[:10]:
            print(f"  - {f[1]}")

asyncio.run(main())
